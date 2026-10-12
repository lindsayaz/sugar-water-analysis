# =============================================================================
# Harvested-area scaling uncertainty propagation for green ET
# =============================================================================
#
# Purpose
# -------
# This script estimates the uncertainty introduced by the harvested-area
# scaling procedure used to adjust country-level green evapotranspiration
# estimates for sugarcane and sugar beet.
#
# The final model result examined here is:
#
#     Global green ET =
#     Sugarcane green ET + Sugar beet green ET
#
# The uncertainty analysis evaluates only the variability introduced by the
# harvested-area scaling procedure. It does not estimate the full uncertainty
# of the underlying crop-water model or other assumptions used elsewhere in
# the study.
#
#
# Method
# ------
# For each country and crop:
#
# 1. The script identifies the harvested-area reference period recorded in
#    the column baseline_used:
#
#       - 1999-2001
#       - 1995-2005
#       - first_nonzero
#       - none
#
# 2. For the 1999-2001 and 1995-2005 baselines, the standard deviation of
#    harvested area is calculated using the annual FAOSTAT harvested-area
#    observations within the relevant reference period.
#
# 3. The coefficient of variation is calculated as:
#
#       CV_i = SD(reference harvested area_i)
#              --------------------------------
#              Mean(reference harvested area_i)
#
#    The column area_ref_final contains the mean reference harvested area
#    that was used in the original scaling procedure.
#
# 4. The harvested-area variability is propagated directly to the scaled
#    green ET estimate:
#
#       u_i = Green ET_scaled_i × CV_i
#
# 5. The country-level uncertainties are combined using the
#    Root-Sum-Square method:
#
#       u_global = sqrt(u_1² + u_2² + ... + u_n²)
#
# 6. Relative uncertainty is calculated as:
#
#       Relative uncertainty (%) =
#       (u_global / Global green ET) × 100
#
#
# Treatment of special cases
# --------------------------
# - baseline_used = "none":
#   No valid harvested-area reference was available. The observation remains
#   part of the reported global green ET total, when a green ET value exists,
#   but it is excluded from the uncertainty propagation.
#
# - baseline_used = "first_nonzero":
#   Only one harvested-area observation was used as the reference. Temporal
#   variability cannot be estimated from a single value, so the harvested-area
#   standard deviation and scaling uncertainty are set to zero.
#
# - Reference periods containing only one valid positive observation:
#   The standard deviation is also set to zero because temporal variability
#   cannot be estimated.
#
#
# Outputs
# -------
# The script prints and exports:
#
# - global sugarcane green ET;
# - global sugar beet green ET;
# - aggregated sugarcane and sugar beet green ET;
# - combined harvested-area scaling uncertainty using RSS;
# - relative harvested-area scaling uncertainty;
# - country-level uncertainty calculations;
# - observations excluded from uncertainty propagation;
# - a global summary table.
#
# =============================================================================


from pathlib import Path

import numpy as np
import pandas as pd


# =============================================================================
# 1. INPUT SETTINGS
# =============================================================================

BASE_DIR = Path(__file__).resolve().parent
input_file = BASE_DIR / "data" / "Country-level blue and green ET.xlsx"

sheet_name = "Sheet1"
target_year = 2021

OUTPUT_DIR = BASE_DIR / "outputs"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
output_file = OUTPUT_DIR / "Harvested-area scaling uncertainty - green ET.xlsx"


# =============================================================================
# 2. CHECK THAT THE INPUT FILE EXISTS
# =============================================================================

if not input_file.exists():
    raise FileNotFoundError(
        "The input Excel file was not found:\n"
        f"{input_file}\n\n"
        "Check that the folder and file names are correct."
    )


# =============================================================================
# 3. READ THE EXCEL FILE
# =============================================================================

df = pd.read_excel(
    input_file,
    sheet_name=sheet_name,
)


# Remove accidental spaces from column names
df.columns = df.columns.astype(str).str.strip()


# =============================================================================
# 4. CHECK REQUIRED COLUMNS
# =============================================================================

required_columns = [
    "iso3",
    "country",
    "year",
    "crop",
    "area_ha",
    "area_ref_final",
    "baseline_used",
    "green_total_m3_scaled",
]

missing_columns = [
    column
    for column in required_columns
    if column not in df.columns
]

if missing_columns:
    raise ValueError(
        "The following required columns are missing from the workbook:\n"
        f"{missing_columns}"
    )


# =============================================================================
# 5. CLEAN AND STANDARDIZE THE DATA
# =============================================================================

numeric_columns = [
    "year",
    "area_ha",
    "area_ref_final",
    "green_total_m3_scaled",
]

for column in numeric_columns:
    df[column] = pd.to_numeric(
        df[column],
        errors="coerce",
    )


# Standardize text columns
for column in ["iso3", "country", "crop"]:
    df[column] = (
        df[column]
        .astype("string")
        .str.strip()
    )


# Standardize baseline labels
df["baseline_used"] = (
    df["baseline_used"]
    .astype("string")
    .str.strip()
    .str.lower()
    .str.replace("–", "-", regex=False)
    .str.replace("—", "-", regex=False)
    .str.replace("_", "-", regex=False)
    .str.replace(" ", "", regex=False)
)


# Standardize common baseline variants
baseline_replacements = {
    "1999to2001": "1999-2001",
    "1995to2005": "1995-2005",
    "first-nonzero": "first_nonzero",
    "firstnonzero": "first_nonzero",
    "first-non-zero": "first_nonzero",
    "na": "none",
    "n/a": "none",
    "nan": "none",
    "<na>": "none",
    "": "none",
}

df["baseline_used"] = df["baseline_used"].replace(
    baseline_replacements
)


# Convert scaled green ET from m³ to km³
df["green_et_km3_scaled"] = (
    df["green_total_m3_scaled"] / 1_000_000_000
)


# Create a standardized crop label for crop-specific summaries
df["crop_standardized"] = (
    df["crop"]
    .astype("string")
    .str.strip()
    .str.lower()
    .str.replace("_", " ", regex=False)
    .str.replace("-", " ", regex=False)
    .str.replace(r"\s+", " ", regex=True)
)


crop_replacements = {
    "sugar cane": "sugarcane",
    "sugarcane": "sugarcane",
    "sc": "sugarcane",
    "sugar beet": "sugar beet",
    "sugarbeet": "sugar beet",
    "sb": "sugar beet",
}

df["crop_standardized"] = df["crop_standardized"].replace(
    crop_replacements
)


# =============================================================================
# 6. VALIDATE THE BASELINE LABELS
# =============================================================================

allowed_baselines = {
    "1999-2001",
    "1995-2005",
    "first_nonzero",
    "none",
}

observed_baselines = set(
    df["baseline_used"]
    .dropna()
    .unique()
)

unknown_baselines = sorted(
    observed_baselines - allowed_baselines
)

if unknown_baselines:
    raise ValueError(
        "Unexpected values were found in baseline_used:\n"
        f"{unknown_baselines}\n\n"
        "Review these labels in the input workbook or add them to the "
        "standardization section of the script."
    )


# =============================================================================
# 7. FUNCTION TO CALCULATE REFERENCE-AREA STATISTICS
# =============================================================================

def calculate_reference_statistics(group: pd.DataFrame) -> pd.Series:
    """
    Calculate harvested-area statistics for one country-crop group.

    Parameters
    ----------
    group
        All annual observations for one country-crop combination.

    Returns
    -------
    pandas.Series
        Baseline label, number of valid reference observations,
        calculated reference-period mean, and reference-period SD.
    """

    baseline_values = (
        group["baseline_used"]
        .dropna()
        .unique()
    )

    if len(baseline_values) == 0:
        baseline = "none"
    elif len(baseline_values) == 1:
        baseline = baseline_values[0]
    else:
        raise ValueError(
            "More than one baseline_used value was found for:\n"
            f"ISO3: {group['iso3'].iloc[0]}\n"
            f"Country: {group['country'].iloc[0]}\n"
            f"Crop: {group['crop'].iloc[0]}\n"
            f"Values: {baseline_values.tolist()}"
        )

    if baseline == "1999-2001":
        reference_values = group.loc[
            group["year"].between(1999, 2001),
            "area_ha",
        ]

    elif baseline == "1995-2005":
        reference_values = group.loc[
            group["year"].between(1995, 2005),
            "area_ha",
        ]

    elif baseline == "first_nonzero":
        positive_values = (
            group.loc[
                group["area_ha"] > 0,
                ["year", "area_ha"],
            ]
            .dropna()
            .sort_values("year")
        )

        if positive_values.empty:
            return pd.Series(
                {
                    "baseline_used_check": baseline,
                    "reference_observation_count": 0,
                    "area_reference_mean_calculated": np.nan,
                    "area_reference_sd": np.nan,
                }
            )

        first_value = float(
            positive_values.iloc[0]["area_ha"]
        )

        return pd.Series(
            {
                "baseline_used_check": baseline,
                "reference_observation_count": 1,
                "area_reference_mean_calculated": first_value,
                "area_reference_sd": 0.0,
            }
        )

    elif baseline == "none":
        return pd.Series(
            {
                "baseline_used_check": baseline,
                "reference_observation_count": 0,
                "area_reference_mean_calculated": np.nan,
                "area_reference_sd": np.nan,
            }
        )

    else:
        raise ValueError(
            f"Unknown baseline_used value: {baseline}"
        )

    reference_values = pd.to_numeric(
        reference_values,
        errors="coerce",
    )

    reference_values = reference_values.dropna()

    # Retain only positive harvested-area observations
    reference_values = reference_values[
        reference_values > 0
    ]

    observation_count = len(reference_values)

    if observation_count == 0:
        reference_mean = np.nan
        reference_sd = np.nan

    elif observation_count == 1:
        reference_mean = float(
            reference_values.iloc[0]
        )
        reference_sd = 0.0

    else:
        reference_mean = float(
            reference_values.mean()
        )

        # Sample standard deviation
        reference_sd = float(
            reference_values.std(ddof=1)
        )

    return pd.Series(
        {
            "baseline_used_check": baseline,
            "reference_observation_count": observation_count,
            "area_reference_mean_calculated": reference_mean,
            "area_reference_sd": reference_sd,
        }
    )


# =============================================================================
# 8. CALCULATE REFERENCE STATISTICS FOR EACH COUNTRY AND CROP
# =============================================================================

reference_statistics_records = []


group_columns = [
    "iso3",
    "country",
    "crop",
    "crop_standardized",
]


for group_keys, group in df.groupby(
    group_columns,
    dropna=False,
    sort=False,
):
    iso3, country, crop, crop_standardized = group_keys

    statistics = calculate_reference_statistics(
        group
    )

    reference_statistics_records.append(
        {
            "iso3": iso3,
            "country": country,
            "crop": crop,
            "crop_standardized": crop_standardized,
            "baseline_used_check": statistics[
                "baseline_used_check"
            ],
            "reference_observation_count": statistics[
                "reference_observation_count"
            ],
            "area_reference_mean_calculated": statistics[
                "area_reference_mean_calculated"
            ],
            "area_reference_sd": statistics[
                "area_reference_sd"
            ],
        }
    )


reference_statistics = pd.DataFrame(
    reference_statistics_records
)


# =============================================================================
# 9. SELECT THE TARGET-YEAR OBSERVATIONS
# =============================================================================

target_df = df.loc[
    df["year"] == target_year,
    [
        "iso3",
        "country",
        "crop",
        "crop_standardized",
        "year",
        "baseline_used",
        "area_ref_final",
        "green_total_m3_scaled",
        "green_et_km3_scaled",
    ],
].copy()


if target_df.empty:
    raise ValueError(
        f"No observations were found for the target year {target_year}."
    )


# Check for duplicate target-year country-crop observations
duplicate_mask = target_df.duplicated(
    subset=[
        "iso3",
        "country",
        "crop",
    ],
    keep=False,
)

if duplicate_mask.any():
    duplicate_rows = target_df.loc[
        duplicate_mask,
        [
            "iso3",
            "country",
            "crop",
            "year",
        ],
    ]

    raise ValueError(
        "Duplicate target-year observations were found for some "
        "country-crop combinations:\n\n"
        f"{duplicate_rows.to_string(index=False)}"
    )


# Merge reference-period statistics into the target-year table
target_df = target_df.merge(
    reference_statistics,
    on=[
        "iso3",
        "country",
        "crop",
        "crop_standardized",
    ],
    how="left",
    validate="one_to_one",
)


# =============================================================================
# 10. COMPARE THE STORED AND RECALCULATED REFERENCE MEANS
# =============================================================================

target_df["reference_mean_difference_ha"] = (
    target_df["area_ref_final"]
    - target_df["area_reference_mean_calculated"]
)


target_df["reference_mean_relative_difference_percent"] = np.where(
    target_df["area_ref_final"].notna()
    & (target_df["area_ref_final"] != 0)
    & target_df["area_reference_mean_calculated"].notna(),
    (
        target_df["reference_mean_difference_ha"]
        / target_df["area_ref_final"]
        * 100
    ),
    np.nan,
)


# =============================================================================
# 11. IDENTIFY OBSERVATIONS INCLUDED IN UNCERTAINTY PROPAGATION
# =============================================================================

target_df["included_in_uncertainty_calculation"] = (
    target_df["baseline_used"].isin(
        [
            "1999-2001",
            "1995-2005",
            "first_nonzero",
        ]
    )
    & target_df["area_ref_final"].notna()
    & (target_df["area_ref_final"] > 0)
    & target_df["area_reference_sd"].notna()
    & target_df["green_et_km3_scaled"].notna()
)


# Record a reason for every excluded observation
def exclusion_reason(row: pd.Series) -> str:
    """Return the principal reason an observation was excluded."""

    if row["baseline_used"] == "none":
        return "No valid harvested-area reference baseline"

    if pd.isna(row["area_ref_final"]):
        return "Missing final reference harvested area"

    if row["area_ref_final"] <= 0:
        return "Final reference harvested area is zero or negative"

    if pd.isna(row["area_reference_sd"]):
        return "Reference-period harvested-area SD could not be calculated"

    if pd.isna(row["green_et_km3_scaled"]):
        return "Missing scaled green ET"

    return "Included"


target_df["uncertainty_calculation_status"] = target_df.apply(
    exclusion_reason,
    axis=1,
)


valid_df = target_df.loc[
    target_df["included_in_uncertainty_calculation"]
].copy()


excluded_df = target_df.loc[
    ~target_df["included_in_uncertainty_calculation"]
].copy()


if valid_df.empty:
    raise ValueError(
        "No valid country-crop observations remained for the "
        "uncertainty calculation."
    )


# =============================================================================
# 12. CALCULATE THE COUNTRY-LEVEL COEFFICIENT OF VARIATION
# =============================================================================
#
# CV_i = SD(reference area_i) / mean(reference area_i)
#
# area_ref_final is the reference-area mean used in the original scaling.

valid_df["area_reference_cv"] = (
    valid_df["area_reference_sd"]
    / valid_df["area_ref_final"]
)


valid_df["area_reference_cv_percent"] = (
    valid_df["area_reference_cv"] * 100
)


# =============================================================================
# 13. PROPAGATE HARVESTED-AREA VARIABILITY TO GREEN ET
# =============================================================================
#
# u_i = Green ET_scaled_i × CV_i

valid_df["green_et_uncertainty_km3"] = (
    valid_df["green_et_km3_scaled"]
    * valid_df["area_reference_cv"]
)


# Add calculated values back into the complete target-year table
target_df = target_df.merge(
    valid_df[
        [
            "iso3",
            "country",
            "crop",
            "area_reference_cv",
            "area_reference_cv_percent",
            "green_et_uncertainty_km3",
        ]
    ],
    on=[
        "iso3",
        "country",
        "crop",
    ],
    how="left",
    validate="one_to_one",
)


# =============================================================================
# 14. CALCULATE GLOBAL GREEN ET RESULTS
# =============================================================================
#
# The central global green ET estimate includes all target-year observations
# with a valid scaled green ET value, including observations that could not
# contribute to the uncertainty propagation.

sugarcane_green_et = target_df.loc[
    target_df["crop_standardized"] == "sugarcane",
    "green_et_km3_scaled",
].sum(
    min_count=1
)


sugar_beet_green_et = target_df.loc[
    target_df["crop_standardized"] == "sugar beet",
    "green_et_km3_scaled",
].sum(
    min_count=1
)


global_green_et = target_df[
    "green_et_km3_scaled"
].sum(
    min_count=1
)


# Confirm that only the two expected crops contribute to the result
unexpected_crops = sorted(
    target_df.loc[
        ~target_df["crop_standardized"].isin(
            [
                "sugarcane",
                "sugar beet",
            ]
        ),
        "crop_standardized",
    ]
    .dropna()
    .unique()
)

if unexpected_crops:
    raise ValueError(
        "Unexpected crop labels were found in the target-year data:\n"
        f"{unexpected_crops}\n\n"
        "The global total should contain only sugarcane and sugar beet."
    )


# =============================================================================
# 15. COMBINE COUNTRY-LEVEL GREEN ET UNCERTAINTIES USING RSS
# =============================================================================
#
# u_global = sqrt(sum(u_i²))

global_green_et_uncertainty_rss = float(
    np.sqrt(
        np.square(
            valid_df["green_et_uncertainty_km3"]
        ).sum()
    )
)


# =============================================================================
# 16. CALCULATE RELATIVE GREEN ET UNCERTAINTY
# =============================================================================

if pd.isna(global_green_et) or global_green_et <= 0:
    raise ValueError(
        "The calculated global green ET is missing, zero, or negative."
    )


relative_uncertainty_fraction = (
    global_green_et_uncertainty_rss
    / global_green_et
)


relative_uncertainty_percent = (
    relative_uncertainty_fraction * 100
)


# =============================================================================
# 17. CALCULATE COVERAGE OF THE UNCERTAINTY ANALYSIS
# =============================================================================

green_et_included_in_uncertainty = valid_df[
    "green_et_km3_scaled"
].sum(
    min_count=1
)


green_et_excluded_from_uncertainty = excluded_df[
    "green_et_km3_scaled"
].sum(
    min_count=1
)


uncertainty_coverage_percent = (
    green_et_included_in_uncertainty
    / global_green_et
    * 100
)


# =============================================================================
# 18. PRINT RESULTS
# =============================================================================

print()
print("=" * 76)
print("HARVESTED-AREA SCALING UNCERTAINTY FOR GLOBAL GREEN ET")
print("=" * 76)

print(f"Target year: {target_year}")

print(
    "Target-year country-crop observations: "
    f"{len(target_df)}"
)

print(
    "Observations included in uncertainty propagation: "
    f"{len(valid_df)}"
)

print(
    "Observations excluded from uncertainty propagation: "
    f"{len(excluded_df)}"
)

print(
    "Share of global green ET covered by uncertainty propagation: "
    f"{uncertainty_coverage_percent:.2f}%"
)

print("-" * 76)

print(
    "Global sugarcane green ET: "
    f"{sugarcane_green_et:.6f} km³"
)

print(
    "Global sugar beet green ET: "
    f"{sugar_beet_green_et:.6f} km³"
)

print(
    "Aggregated green ET: sugarcane + sugar beet: "
    f"{global_green_et:.6f} km³"
)

print("-" * 76)

print(
    "Combined harvested-area scaling uncertainty for green ET (RSS): "
    f"±{global_green_et_uncertainty_rss:.6f} km³"
)

print(
    "Relative harvested-area scaling uncertainty for green ET: "
    f"{relative_uncertainty_percent:.4f}%"
)

print("-" * 76)

print(
    "Final reported result: "
    f"{global_green_et:.3f} ± "
    f"{global_green_et_uncertainty_rss:.3f} km³ "
    f"({relative_uncertainty_percent:.2f}%)"
)

print("=" * 76)


# =============================================================================
# 19. CREATE THE GLOBAL SUMMARY TABLE
# =============================================================================

summary = pd.DataFrame(
    {
        "Parameter": [
            "Target year",
            "Target-year country-crop observations",
            "Observations included in uncertainty propagation",
            "Observations excluded from uncertainty propagation",
            "Share of global green ET covered by uncertainty propagation",
            "Global sugarcane green ET",
            "Global sugar beet green ET",
            "Aggregated green ET: sugarcane + sugar beet",
            "Combined harvested-area scaling uncertainty using RSS",
            "Relative harvested-area scaling uncertainty",
            "Final result - central estimate",
            "Final result - uncertainty",
        ],
        "Value": [
            target_year,
            len(target_df),
            len(valid_df),
            len(excluded_df),
            uncertainty_coverage_percent,
            sugarcane_green_et,
            sugar_beet_green_et,
            global_green_et,
            global_green_et_uncertainty_rss,
            relative_uncertainty_percent,
            global_green_et,
            global_green_et_uncertainty_rss,
        ],
        "Unit": [
            "year",
            "observations",
            "observations",
            "observations",
            "%",
            "km³",
            "km³",
            "km³",
            "km³",
            "%",
            "km³",
            "km³",
        ],
        "Description": [
            "Year used for the final green ET estimate.",
            "Number of country-crop records available for the target year.",
            "Records with sufficient information for uncertainty propagation.",
            "Records lacking a valid reference area, SD, or scaled green ET.",
            "Percentage of the global green ET represented by records included in the uncertainty propagation.",
            "Rainfall-derived consumptive water use for sugarcane, summed across countries.",
            "Rainfall-derived consumptive water use for sugar beet, summed across countries.",
            "Global rainfall-derived consumptive water use for sugarcane and sugar beet.",
            "Country-level green ET uncertainties combined using Root-Sum-Square.",
            "RSS uncertainty divided by aggregated global green ET.",
            "Aggregated global green ET estimate.",
            "Harvested-area scaling uncertainty associated with the global green ET estimate.",
        ],
    }
)


# =============================================================================
# 20. CREATE AN EXCLUSION SUMMARY
# =============================================================================

if excluded_df.empty:
    exclusion_summary = pd.DataFrame(
        {
            "Exclusion reason": [],
            "Number of observations": [],
            "Excluded green ET (km³)": [],
        }
    )

else:
    exclusion_summary = (
        excluded_df.groupby(
            "uncertainty_calculation_status",
            dropna=False,
        )
        .agg(
            number_of_observations=(
                "country",
                "size",
            ),
            excluded_green_et_km3=(
                "green_et_km3_scaled",
                "sum",
            ),
        )
        .reset_index()
        .rename(
            columns={
                "uncertainty_calculation_status": (
                    "Exclusion reason"
                ),
                "number_of_observations": (
                    "Number of observations"
                ),
                "excluded_green_et_km3": (
                    "Excluded green ET (km³)"
                ),
            }
        )
    )


# =============================================================================
# 21. EXPORT RESULTS TO EXCEL
# =============================================================================

with pd.ExcelWriter(
    output_file,
    engine="openpyxl",
) as writer:

    summary.to_excel(
        writer,
        sheet_name="Global results",
        index=False,
    )

    valid_df.to_excel(
        writer,
        sheet_name="Country calculations",
        index=False,
    )

    excluded_df.to_excel(
        writer,
        sheet_name="Excluded observations",
        index=False,
    )

    exclusion_summary.to_excel(
        writer,
        sheet_name="Exclusion summary",
        index=False,
    )

    reference_statistics.to_excel(
        writer,
        sheet_name="Reference statistics",
        index=False,
    )

    target_df.to_excel(
        writer,
        sheet_name="All target-year records",
        index=False,
    )


# =============================================================================
# 22. APPLY BASIC FORMATTING TO THE OUTPUT WORKBOOK
# =============================================================================

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


workbook = load_workbook(output_file)

header_fill = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7",
)

header_font = Font(
    bold=True,
)

for worksheet in workbook.worksheets:

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter

        for cell in column_cells:
            value = cell.value

            if value is not None:
                max_length = max(
                    max_length,
                    len(str(value)),
                )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 12),
            45,
        )


workbook.save(output_file)


print()
print("Results were saved to:")
print(output_file)
